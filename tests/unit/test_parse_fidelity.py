"""P2 解析保真度单测：真实样例文档的关键事实必须能被解析器抽取。"""
from __future__ import annotations

import io
from pathlib import Path

from app.adapters.parser.dispatcher import parse_file

SAMPLE = Path(__file__).resolve().parent.parent / "eval" / "sample_data"


def _parse(path: Path):
    return parse_file(io.BytesIO(path.read_bytes()), path.name)


def test_employee_handbook_facts_extracted():
    doc = _parse(SAMPLE / "员工手册.md")
    for fact in ("入职满 1 年：年假 5 天", "最高累计 20 天", "5 个工作日内打款", "150 元/天"):
        assert fact in doc.text, f"解析丢失事实：{fact}"
    assert doc.metadata["source"] == "员工手册.md"


def test_travel_policy_facts_extracted():
    doc = _parse(SAMPLE / "docs" / "差旅政策.md")
    for fact in ("500 元/晚", "提前 7 天", "每公里 0.8 元", "3 个工作日内需提交出差总结"):
        assert fact in doc.text, f"解析丢失事实：{fact}"


def test_docx_table_cells_extracted():
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    d = DocxDocument()
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "住宿标准"
    table.cell(1, 0).text = "一线城市"
    table.cell(1, 1).text = "500 元/晚"
    d.save(buf)
    doc = parse_file(io.BytesIO(buf.getvalue()), "标准.docx")
    assert "一线城市" in doc.text and "500 元/晚" in doc.text and "住宿标准" in doc.text
