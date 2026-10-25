"""标题硬边界 / 上下文前缀 / 页码溯源单测（P1-5，DEVELOPMENT.md 5.2）。

覆盖：标题强制断块、同 section 不误断、多页 page_spans 映射、前缀模板渲染、
无 section 降级、超长前缀截断、include_page、enabled=false 回归、pptx slide 页码、
表格块页码。
"""
from __future__ import annotations

import io

from app.adapters.chunker import split_chunks
from app.adapters.parser import parse_file
from app.core.domain import Document

HB = {"enabled": True}
CP = {"enabled": True, "template": "{doc_title} > {section}", "include_page": False}
CP_PAGE = {"enabled": True, "template": "{doc_title} > {section}", "include_page": True}


def _doc(text: str, **meta) -> Document:
    metadata = {"source": "t.md", "kind": "md", **meta}
    return Document(document_id="d1", filename="t.md", title="员工手册",
                    content_hash="h", text=text, metadata=metadata)


def _parents(chunks):
    return [c for c in chunks if c.parent_id is None]


# -- 标题硬边界 ------------------------------------------------------------------------

def test_hard_boundary_flushes_on_section_change():
    text = "第一章的内容行。\n第二章的内容行。\n第二章的第二行。"
    sections = [("第一章", 0), ("第二章", 1)]
    doc = _doc(text, sections=sections)
    without = _parents(split_chunks(doc))
    with_hb = _parents(split_chunks(doc, hard_boundary=HB))
    assert len(without) == 1, "关闭硬边界时小文本聚为一块（旧行为）"
    assert len(with_hb) == 2, "硬边界必须按 section 变更断块"
    assert with_hb[0].metadata["section"] == "第一章"
    assert with_hb[1].text.startswith("第二章的内容行"), "新块以标题后首行为起点"


def test_hard_boundary_no_false_break_within_section():
    text = "第一章第一行。\n第一章第二行。\n第一章第三行。"
    doc = _doc(text, sections=[("第一章", 0)])
    parents = _parents(split_chunks(doc, hard_boundary=HB))
    assert len(parents) == 1, "同 section 连续内容不得被硬边界误断"


def test_table_atom_flushes_and_keeps_order():
    """表格原子在两个 section 之间：硬边界下不与正文粘块，顺序保持正文→表→正文。"""
    text = ("第一章正文。\n[TABLE]\n| A |\n|---|\n| 1 |\n[/TABLE]\n第二章正文。")
    doc = _doc(text, sections=[("第一章", 0), ("第二章", 6)])
    chunks = split_chunks(doc, hard_boundary=HB, table_aware={"enabled": True, "max_rows_per_chunk": 3, "json_max_bytes": 4096})
    texts = [c.text for c in chunks]
    body1 = next(i for i, t in enumerate(texts) if "第一章正文" in t)
    table = next(i for i, t in enumerate(texts) if "[TABLE]" in t)
    body2 = next(i for i, t in enumerate(texts) if "第二章正文" in t)
    assert body1 < table < body2


# -- 页码溯源 --------------------------------------------------------------------------

def test_page_resolved_from_page_spans():
    text = "第一页内容。\n第一页内容续。\n第二页内容。\n第二页内容续。"
    doc = _doc(text, page_spans=[[0, 1, 1], [2, 3, 2]])
    # 页码不触发断块（规格：仅 section 变更断块）——缩小 parent 目标让每行独立成块验证映射
    parents = _parents(split_chunks(doc, parent_target_tokens=5, child_target_tokens=3))
    assert [p.metadata.get("page") for p in parents] == [1, 1, 2, 2]
    assert "page_spans" not in parents[0].metadata, "page_spans 是文档级映射，不随 chunk 落库"


def test_page_absent_when_no_spans():
    doc = _doc("普通文本内容。")
    parents = _parents(split_chunks(doc))
    assert all("page" not in p.metadata for p in parents)


def test_table_chunk_carries_page():
    text = "第一页正文。\n[TABLE]\n| A |\n|---|\n| 1 |\n[/TABLE]"
    doc = _doc(text, page_spans=[[0, 0, 1], [1, 4, 2]])
    chunks = split_chunks(doc, table_aware={"enabled": True, "max_rows_per_chunk": 3,
                                            "json_max_bytes": 4096})
    table_chunk = next(c for c in chunks if c.metadata.get("chunk_type") == "table")
    assert table_chunk.metadata.get("page") == 2, "表格块按起始行查 page_spans"


# -- 上下文前缀 ------------------------------------------------------------------------

def test_prefix_rendered_for_all_chunks():
    text = "第一章正文。\n第一章正文续。"
    doc = _doc(text, sections=[("第一章", 0)])
    chunks = split_chunks(doc, context_prefix=CP)
    assert chunks, "应产出 chunk"
    for c in chunks:
        assert c.text.startswith("员工手册 > 第一章\n"), f"父子块一致加前缀：{c.text[:30]!r}"


def test_prefix_fallback_without_section():
    doc = _doc("无章节归属的文本。")
    chunks = split_chunks(doc, context_prefix=CP)
    assert chunks[0].text.startswith("员工手册\n"), "无 section 降级为仅 doc_title"


def test_prefix_include_page():
    text = "第一页内容。\n第一页内容续。"
    doc = _doc(text, page_spans=[[0, 1, 1]])
    chunks = split_chunks(doc, context_prefix=CP_PAGE)
    assert chunks[0].text.startswith("员工手册 (p1)\n")


def test_prefix_truncated_over_64_tokens():
    long_title = "超长文档标题" * 40     # ≈240 字符 ≫ 64 token
    doc = Document(document_id="d1", filename="t.md", title=long_title,
                   content_hash="h", text="正文内容。", metadata={"source": "t.md", "kind": "md"})
    chunks = split_chunks(doc, context_prefix=CP)
    prefix = chunks[0].text.split("\n")[0]
    assert len(prefix) <= 64 * 2, "超长前缀必须截断（64 token × ratio≈2）"


def test_prefix_disabled_is_legacy():
    doc = _doc("正文内容。", sections=[("第一章", 0)])
    chunks = split_chunks(doc)
    assert chunks[0].text == "正文内容。", "关闭前缀时行为与旧版一致"


# -- pptx slide 页码 --------------------------------------------------------------------

def test_pptx_page_spans_and_chunk_page():
    from pptx import Presentation

    buf = io.BytesIO()
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "第一页标题"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "第二页标题"
    prs.save(buf)
    parsed = parse_file(io.BytesIO(buf.getvalue()), "t.pptx")
    assert parsed.metadata["page_spans"] == [[0, 0, 1], [1, 1, 2]], "pptx 以 slide 号为页码"
    # 页码不触发断块（规格：仅 section 变更断块）——缩小 parent 目标逐行验证映射
    pages = [p.metadata.get("page") for p in _parents(split_chunks(parsed, parent_target_tokens=5, child_target_tokens=3))]
    assert pages == [1, 2]


# -- Markdown 管道表格（无 [TABLE] 标记，如 .md 文件） ---------------------------------

def test_md_pipe_table_is_atomic_with_header():
    """md 表格没有解析器元数据：开启 table_aware 后也必须整体成块且保留表头。"""
    text = ("## 四、常见 OOM 类型\n\n"
            "| 类型 | 报错信息 | 根因 |\n|---|---|---|\n"
            "| 堆 OOM | Java heap space | 对象过多 |\n"
            "| 栈 OOM | StackOverflowError | 递归过深 |\n")
    doc = _doc(text)
    chunks = split_chunks(doc, table_aware={"enabled": True, "max_rows_per_chunk": 30,
                                            "json_max_bytes": 4096})
    table = next(c for c in chunks if c.metadata.get("chunk_type") == "table")
    assert "| 类型 | 报错信息 | 根因 |" in table.text, "表头必须在同一块内"
    assert "堆 OOM" in table.text and "栈 OOM" in table.text
    assert "|---|---|---|" in table.text


def test_md_pipe_table_long_splits_with_header():
    """超长 md 表按行组切分，每组都带完整表头（与 [TABLE] 行为一致）。"""
    rows = "\n".join(f"| 行{i} | {'内容' * 30} |" for i in range(12))
    text = "| 序号 | 说明 |\n|---|---|\n" + rows
    doc = _doc(text)
    chunks = split_chunks(doc, parent_target_tokens=80, table_aware={"enabled": True,
                                                                     "max_rows_per_chunk": 4,
                                                                     "json_max_bytes": 4096})
    tables = [c for c in chunks if c.metadata.get("chunk_type") == "table"]
    assert len(tables) >= 2, "超长表应拆成多块"
    assert all("| 序号 | 说明 |" in c.text for c in tables), "每组必须携带表头"


def test_md_pipe_table_inside_code_fence_not_split():
    """代码围栏里的管道示例不能被当成表格原子块。"""
    text = "```text\n| A | B |\n|---|---|\n| 1 | 2 |\n```\n"
    doc = _doc(text)
    chunks = split_chunks(doc, table_aware={"enabled": True, "max_rows_per_chunk": 30,
                                            "json_max_bytes": 4096})
    assert not any(c.metadata.get("chunk_type") == "table" for c in chunks)
