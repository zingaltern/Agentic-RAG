"""表格一等公民单测（P1-4，DEVELOPMENT.md 5.1/5.2）。

覆盖：小表单块不切断 / 大表行组带表头 / 正文多表混排不粘块 / reindex 对齐与错位丢弃 /
table_json 截断 / docx 保序标记化 / enabled=false 回归 / 探针键剔除 / cleaner 表格豁免。
"""
from __future__ import annotations

import io
import json

from app.adapters.cleaner import DefaultCleaner
from app.adapters.chunker import split_chunks
from app.adapters.parser import parse_file, reindex_tables
from app.core.config import load_config
from app.core.domain import Document

TA = {"enabled": True, "max_rows_per_chunk": 3, "json_max_bytes": 4096}


def _doc(text: str, tables: list[dict] | None = None, **meta) -> Document:
    metadata = {"source": "t.md", "kind": "md", **meta}
    if tables is not None:
        metadata["tables"] = tables
    return Document(document_id="d1", filename="t.md", title="t",
                    content_hash="h", text=text, metadata=metadata)


def _table_md(rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "|".join("---" for _ in rows[0]) + "|"]
    lines.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n".join(lines)


# -- 小表：独立父块不切断 ------------------------------------------------------------

def test_small_table_single_parent_with_json():
    tables = [{"index": 0, "header": ["字段", "说明"], "rows": [["年假", "5 天"]],
               "page": None}]
    text = "正文段落。\n[TABLE]\n" + _table_md(tables[0]["header"] + []) + "\n[/TABLE]"
    # 构造与 tables[0] 一致的 markdown（表头+1 行数据）
    md = _table_md([["字段", "说明"], ["年假", "5 天"]])
    text = "正文段落。\n[TABLE]\n" + md + "\n[/TABLE]\n结尾段落。"
    chunks = split_chunks(_doc(text, tables), parent_target_tokens=512,
                          child_target_tokens=128, table_aware=TA)
    table_chunks = [c for c in chunks if c.metadata.get("chunk_type") == "table"]
    assert len(table_chunks) == 1, "小表必须为独立单父块"
    tc = table_chunks[0]
    assert tc.parent_id is None, "表格块为父块，不产子块"
    payload = json.loads(tc.metadata["table_json"])
    assert payload["header"] == ["字段", "说明"]
    assert payload["rows"] == [["年假", "5 天"]]
    assert "[TABLE]" in tc.text and "年假" in tc.text
    assert tc.metadata.get("table_part") is None


# -- 大表：行组切分强制带表头 --------------------------------------------------------

def test_large_table_row_groups_carry_header():
    header = ["月份", "营收"]
    rows = [[f"2026-{i:02d}", str(i * 100)] for i in range(1, 11)]
    md = _table_md([header] + rows)          # parent_target=30 → 超限；max_rows=3 → 4 组
    tables = [{"index": 0, "header": header, "rows": rows, "page": None}]
    chunks = split_chunks(_doc("[TABLE]\n" + md + "\n[/TABLE]", tables),
                          parent_target_tokens=30, child_target_tokens=10, table_aware=TA)
    table_chunks = [c for c in chunks if c.metadata.get("chunk_type") == "table"]
    assert [c.metadata.get("table_part") for c in table_chunks] == ["1/4", "2/4", "3/4", "4/4"]
    for tc in table_chunks:
        assert tc.text.split("\n")[0].startswith("| 月份"), "每组必须携带完整表头"
        assert tc.text.split("\n")[1] == "|---|---|"
        assert tc.metadata["table_json"], "行组同样携带 table_json"


# -- 正文与多表混排不粘块 ------------------------------------------------------------

def test_text_and_two_tables_not_glued():
    md1 = _table_md([["A"], ["1"]])
    md2 = _table_md([["B"], ["2"]])
    tables = [{"index": 0, "header": ["A"], "rows": [["1"]], "page": None},
              {"index": 1, "header": ["B"], "rows": [["2"]], "page": None}]
    text = "前言内容。\n[TABLE]\n" + md1 + "\n[/TABLE]\n中间内容。\n[TABLE]\n" + md2 + "\n[/TABLE]"
    chunks = split_chunks(_doc(text, tables), table_aware=TA)
    kinds = [(c.metadata.get("chunk_type"), "TABLE" in c.text) for c in chunks]
    assert kinds.count(("table", True)) == 2, "两个表格各自独立成块"
    text_chunks = [c for c in chunks if c.metadata.get("chunk_type") is None]
    assert all("[TABLE]" not in c.text for c in text_chunks), "正文块不得粘入表格"
    assert any("前言" in c.text for c in text_chunks) and any("中间" in c.text for c in text_chunks)
    order = [c.order for c in chunks]
    assert order == sorted(order)
    # 顺序校验：前言 → 表1 → 中间 → 表2
    seq = [c.text for c in sorted(chunks, key=lambda x: x.order)]
    assert seq[0].startswith("前言") and "[TABLE]" in seq[1] \
        and "中间内容" in seq[2] and "[TABLE]" in seq[3]


# -- 探针键剔除（评审附加要求） --------------------------------------------------------

def test_probe_keys_stripped_from_chunks():
    md = _table_md([["A"], ["1"]])
    tables = [{"index": 0, "header": ["A"], "rows": [["1"]], "page": 1}]
    text = "正文。\n[TABLE]\n" + md + "\n[/TABLE]"
    doc = _doc(text, tables, text_stats={"total_chars": 100, "page_count": 1,
                                         "pages_with_text": 1}, sections=[("s", 0)])
    chunks = split_chunks(doc, table_aware=TA)
    assert chunks
    for c in chunks:
        assert "text_stats" not in c.metadata, "解析探针键不得随 chunk 落库"
        assert "tables" not in c.metadata, "tables 列表不得随 chunk 落库（由 table_json 承载）"
        assert "sections" not in c.metadata


# -- table_json 截断 ------------------------------------------------------------------

def test_table_json_truncated():
    huge_cell = "很长的单元格" * 2000
    md = _table_md([["A"], [huge_cell]])
    tables = [{"index": 0, "header": ["A"], "rows": [[huge_cell]], "page": None}]
    ta = {"enabled": True, "max_rows_per_chunk": 3, "json_max_bytes": 512}
    chunks = split_chunks(_doc("[TABLE]\n" + md + "\n[/TABLE]", tables), table_aware=ta)
    tc = next(c for c in chunks if c.metadata.get("chunk_type") == "table")
    tj = tc.metadata["table_json"]
    assert len(tj.encode("utf-8")) <= 512, "table_json 必须按 json_max_bytes 截断"


# -- enabled=false 回归 ----------------------------------------------------------------

def test_disabled_keeps_legacy_behavior():
    md = _table_md([["A"], ["1"]])
    text = "前言。\n[TABLE]\n" + md + "\n[/TABLE]"
    chunks = split_chunks(_doc(text), table_aware=None)
    assert not any(c.metadata.get("chunk_type") == "table" for c in chunks), "关闭时无表格块语义"
    assert any("[TABLE]" in c.text for c in chunks), "关闭时 [TABLE] 行按普通文本聚合（旧行为）"


# -- reindex_tables --------------------------------------------------------------------

def test_reindex_drops_cleaned_tables():
    tables = [{"index": 0, "header": ["A"], "rows": [["1"]], "page": None},
              {"index": 1, "header": ["B"], "rows": [["2"]], "page": None}]
    md1 = _table_md([["A"], ["1"]])
    doc = _doc("x\n[TABLE]\n" + md1 + "\n[/TABLE]", tables)
    reindex_tables(doc)
    assert len(doc.metadata["tables"]) == 1, "被清洗删掉的表 JSON 必须丢弃"
    assert doc.metadata["tables"][0]["header"] == ["A"]
    assert doc.metadata["tables"][0]["index"] == 0, "重对齐后 index 连续"


def test_reindex_noop_when_aligned():
    md = _table_md([["A"], ["1"]])
    tables = [{"index": 0, "header": ["A"], "rows": [["1"]], "page": None}]
    doc = _doc("[TABLE]\n" + md + "\n[/TABLE]", tables)
    reindex_tables(doc)
    assert doc.metadata["tables"] == tables


# -- docx 保序标记化 --------------------------------------------------------------------

def test_docx_table_in_document_order():
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("表格前的段落")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "字段"
    t.rows[0].cells[1].text = "说明"
    t.rows[1].cells[0].text = "年假"
    t.rows[1].cells[1].text = "5 天"
    doc.add_paragraph("表格后的段落")
    doc.save(buf)
    parsed = parse_file(io.BytesIO(buf.getvalue()), "t.docx")
    assert parsed.text.index("表格前的段落") < parsed.text.index("[TABLE]") \
        < parsed.text.index("[/TABLE]") < parsed.text.index("表格后的段落"), "docx 表格必须按文档流保序"
    assert parsed.metadata["tables"][0]["header"] == ["字段", "说明"]
    chunked = split_chunks(parsed, table_aware=TA)
    assert any(c.metadata.get("chunk_type") == "table" for c in chunked)


# -- cleaner 表格豁免 --------------------------------------------------------------------

def test_cleaner_normalize_whitelist_and_quality_exempt():
    cleaner = DefaultCleaner(enabled=True, normalize=True, normalize_punct=True, noise=False,
                             dedup=False, quality=True, min_chunk_chars=20)
    text = ("这是足够长的一段正文内容，包含很多实词和说明文字，确保质量过滤可以通过。\n"
            "[TABLE]\n| 字段　说明 |\n|---|---|\n| 年假　五天 |\n[/TABLE]")
    cleaned = cleaner.clean(text)
    assert "[TABLE]" in cleaned, "含表格段落豁免质量过滤（不会因表外内容被误杀）"
    table_block = cleaned[cleaned.index("[TABLE]"):cleaned.index("[/TABLE]")]
    assert "字段　说明" in table_block, "表格内容全角空格不被 normalize 改写（白名单）"
    assert "这是足够长的一段正文内容,包含很多实词和说明文字,确保质量过滤可以通过." in cleaned, "表外正文仍被 normalize"


def test_cleaner_dedup_exempt_table_paragraph():
    cleaner = DefaultCleaner(enabled=True, normalize=False, noise=False,
                             dedup=True, quality=False)
    body = "完全相同的一段正文内容用来触发近重复去重判定逻辑。" * 3
    text = body + "\n[TABLE]\n" + _table_md([["A"], ["1"]]) + "\n[/TABLE]"
    cleaned = cleaner.clean(text)
    assert "[TABLE]" in cleaned, "含表格段落豁免去重"
