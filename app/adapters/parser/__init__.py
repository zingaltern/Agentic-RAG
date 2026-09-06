"""文档解析（DEVELOPMENT.md 5.1）。

契约（3.9）：`parse_file(file_stream: BinaryIO, filename: str) -> Document`
- 支持 .pdf/.docx/.xlsx/.pptx/.md/.txt，text 归一为 UTF-8。
- metadata 保留键：source（原文件名）、kind（扩展名）、page（int/None）、author（str/None）。
- 解析探针（P0-1）：metadata["text_stats"] = {total_chars, page_count, pages_with_text,
  garbled_ratio}（garbled_ratio 为 U+FFFD 乱码占比，全部解析器统一在 _build_document 计算）。
- 损坏/不支持 → ParseError（code 30003）。
"""
from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path
from typing import BinaryIO
from uuid import uuid4

from app.core.domain import Document
from app.core.errors import ParseError
from app.core.logging import get_logger

logger = get_logger("parser")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".txt",
                       ".html", ".htm", ".eml", ".msg"}


def _build_document(filename: str, raw: bytes, text: str, metadata: dict) -> Document:
    """按 2.3 约定生成 document_id（uuid4 hex 前 12 位）与 content_hash（sha256 原始字节）。

    content_hash 写入 metadata：Chunk 继承后可用于幂等去重检索（W2 5.10）。
    P0-1 探针：garbled_ratio（U+FFFD 乱码占比）在此对全部解析器统一计算——
    txt/html/eml 的 errors="replace" 解码同样产生 U+FFFD，必须纳入质量门禁视野；
    text_stats 缺失时（非 PDF 解析器）按单页文档补齐。
    """
    content_hash = hashlib.sha256(raw).hexdigest()
    metadata = {
        "source": filename,
        "kind": Path(filename).suffix.lstrip(".").lower(),
        "content_hash": content_hash,
        **metadata,
    }
    stats = metadata.setdefault("text_stats", {
        "total_chars": len(text),
        "page_count": 1,
        "pages_with_text": 1 if text.strip() else 0,
    })
    garbled = text.count("\ufffd")
    stats["garbled_ratio"] = (garbled / len(text)) if text else 0.0
    return Document(
        document_id=uuid4().hex[:12],
        filename=filename,
        title=Path(filename).stem,
        content_hash=content_hash,
        text=text,
        metadata=metadata,
    )


def evaluate_parse_quality(text: str, text_stats: dict, *,
                           min_text_chars: int, max_garbled_ratio: float,
                           min_pages_with_text_ratio: float) -> list[str]:
    """P0-1 解析质量检查（模块级纯函数，P3-8 解析器降级链复用）。

    返回 flag 列表（空 = 通过）：
    - empty_text：去首尾空白后文本少于 min_text_chars（含扫描件抽出空文本）；
    - high_garbled：U+FFFD 乱码占比超过 max_garbled_ratio（text_stats 缺
      garbled_ratio 时按当前文本现算兜底）；
    - few_pages_with_text：仅对 page_count>1 的多页 PDF 生效——有实质文本页占比
      pages_with_text/page_count 低于 min_pages_with_text_ratio（疑似扫描件大部分页无文本层）。
    """
    flags: list[str] = []
    stats = text_stats or {}
    if len((text or "").strip()) < min_text_chars:
        flags.append("empty_text")
    ratio = stats.get("garbled_ratio")
    if ratio is None:
        total = len(text or "")
        ratio = (text or "").count("\ufffd") / total if total else 0.0
    if ratio > max_garbled_ratio:
        flags.append("high_garbled")
    page_count = int(stats.get("page_count") or 0)
    pages_with_text = int(stats.get("pages_with_text") or 0)
    if page_count > 1 and (pages_with_text / page_count) < min_pages_with_text_ratio:
        flags.append("few_pages_with_text")
    return flags


_TABLE_BLOCK_RE = re.compile(r"\[TABLE\]\n(.*?)\n\[/TABLE\]", flags=re.S)


def _md_header_cells(block: str) -> list[str]:
    """从 [TABLE] 块提取表头单元格（第一条非分隔行的管道单元格）。"""
    for line in block.split("\n"):
        s = line.strip()
        if not s.startswith("|"):
            continue
        if set(s) <= {"|", "-", " ", ":"}:
            continue   # |---|---| 分隔行
        return [c.strip() for c in s.strip("|").split("|")]
    return []


def reindex_tables(doc: Document) -> None:
    """P1-4：清洗后按 [TABLE] 出现顺序重对齐 metadata["tables"]（DEVELOPMENT 5.1）。

    - 数量一致 → 无操作（normalize 白名单保证表内容未被改写，顺序天然对齐）；
    - cleaner（noise/dedup/quality）删掉了含表格的段落 → 按表头单元格匹配重对齐，
      对不上号的 JSON 丢弃并 warning（表格正文还在 text 中、JSON 缺失时降级为纯文本检索）。
    """
    tables = doc.metadata.get("tables")
    if not tables:
        return
    blocks = _TABLE_BLOCK_RE.findall(doc.text or "")
    if len(blocks) == len(tables):
        return
    kept: list[dict] = []
    used: set[int] = set()
    for block in blocks:
        header = _md_header_cells(block)
        match = next((i for i, t in enumerate(tables)
                      if i not in used
                      and [str(c) for c in (t.get("header") or [])] == header), None)
        if match is None:   # 表头匹配失败（空表头等）→ 取剩余未用表的顺序兜底
            match = next((i for i in range(len(tables)) if i not in used), None)
        if match is None:
            break
        used.add(match)
        kept.append({**tables[match], "index": len(kept)})
    if len(kept) != len(tables):
        logger.warning(
            f"清洗后表格 JSON 重对齐：{len(tables)}→{len(kept)}，"
            f"丢弃 {len(tables) - len(kept)} 个无对应 [TABLE] 块的登记",
            extra={"extra_fields": {"before": len(tables), "after": len(kept),
                                    "filename": doc.filename}},
        )
    doc.metadata["tables"] = kept


def _check_magic(ext: str, raw: bytes, filename: str, config: object | None = None) -> None:
    """S2（V5.1）：上传前按魔数校验真实格式与扩展名一致（pdf %PDF；zip 系 PK\x03\x04；txt/md 文本嗅探）。

    不一致抛 ParseError（伪造扩展名/损坏文件统一语义）。
    """
    if ext == ".pdf":
        if not raw.startswith(b"%PDF"):
            raise ParseError(f"文件格式与扩展名不符（非 PDF）: {filename!r}")
    elif ext in (".docx", ".xlsx", ".pptx"):
        if not raw.startswith(b"PK\x03\x04"):
            raise ParseError(f"文件格式与扩展名不符（非 Office zip）: {filename!r}")
        _check_zip_bomb(raw, filename, config)
    else:  # .md / .txt
        sample = raw[:512]
        if sample and b"\x00" in sample:
            raise ParseError(f"文件格式与扩展名不符（疑似二进制冒充文本）: {filename!r}")


def _check_zip_bomb(raw: bytes, filename: str, config: object | None = None) -> None:
    """S2（V5.1）：zip 解压炸弹防御 —— 解压后总大小/条目数上限（默认 200MB/1000 条）。"""
    import io as _io
    import zipfile as _zipfile

    total_mb = int(getattr(config, "zip_max_total_mb", 200)) if config is not None else 200
    max_entries = int(getattr(config, "zip_max_entries", 1000)) if config is not None else 1000
    try:
        with _zipfile.ZipFile(_io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            if len(infos) > max_entries:
                raise ParseError(f"zip 条目数超过上限（{max_entries}）: {filename!r}")
            total = sum(info.file_size for info in infos)
            if total > total_mb * 1024 * 1024:
                raise ParseError(f"zip 解压总大小超过上限（{total_mb}MB）: {filename!r}")
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"zip 校验失败: {filename!r}（{exc}）") from exc


def parse_file(file_stream: BinaryIO, filename: str, config: object | None = None) -> Document:
    """按扩展名分发解析；统一异常包装为 ParseError（5.1）；S2 魔数/zip 炸弹预检。"""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"不支持的文档格式: {filename!r}（支持：{sorted(e for e in SUPPORTED_EXTENSIONS)}）")
    try:
        raw = file_stream.read()
    except Exception as exc:
        raise ParseError(f"读取文件失败: {filename}") from exc
    if not isinstance(raw, bytes):
        raw = bytes(raw)
    _check_magic(ext, raw, filename, config)
    try:
        parser = PARSERS.get(ext)
        if parser is None:
            raise ParseError(f"不支持的文档格式: {filename!r}")
        text, meta = parser(raw)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"{filename} 解析失败: {exc}") from exc
    return _build_document(filename, raw, text or "", meta)


def _extract_markdown_sections(text: str) -> list[tuple[str, int]]:
    """提取 Markdown 章节标题（T5 元数据富化）：返回 [(title, 行号)]。"""
    sections: list[tuple[str, int]] = []
    for line_no, line in enumerate(text.split("\n")):
        m = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", line)
        if m:
            sections.append((m.group(1).strip(), line_no))
    return sections


def _normalize_pdf_date(raw: str | None) -> str | None:
    """PDF 日期（D:YYYYMMDD...）→ YYYY-MM-DD；无法解析返回 None。"""
    if not raw:
        return None
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", raw)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None


def _pdf_body_size(pdf) -> float:
    """PDF 主体字号估计：全部 span 字号的中位数（标题字号显著大于主体）。"""
    import statistics

    sizes: list[float] = []
    for page in pdf:
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                sizes.extend(span["size"] for span in line.get("spans", []))
    return statistics.median(sizes) if sizes else 12.0


def _pdf_page_headings(page, body_size: float) -> list[str]:
    """按字号启发式提取页内标题行（字号 ≥ 主体 + 2pt）。"""
    headings: list[str] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            sizes = [span["size"] for span in line.get("spans", [])]
            if not sizes:
                continue
            if max(sizes) >= body_size + 2.0:
                title = "".join(span["text"] for span in line.get("spans", [])).strip()
                if title:
                    headings.append(title)
    return headings


def _parse_text(raw: bytes) -> tuple[str, dict]:
    """txt/md：按 UTF-8 解码，非 UTF-8 字节以替换字符归一（5.1）。"""
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    meta: dict = {"page": None, "author": None}
    sections = _extract_markdown_sections(text)
    if sections:
        meta["sections"] = sections        # T5：章节标题（title, 行号）
    return text, meta


def _parse_pdf(raw: bytes) -> tuple[str, dict]:
    """PDF：pymupdf 逐页抽取文本，表格经 find_tables 转 Markdown 表格按阅读顺序合并（保真）。"""
    import pymupdf

    try:
        pdf = pymupdf.open(stream=raw, filetype="pdf")
    except Exception as exc:
        raise ParseError(f"PDF 解析失败（文件损坏或非 PDF）: {exc}") from exc
    try:
        body_size = _pdf_body_size(pdf)
        parts: list[str] = []
        sections: list[tuple[str, int]] = []
        global_line = 0
        # P0-1 探针：逐页统计（total_chars / pages_with_text），供解析质量门禁识别扫描件
        total_chars = 0
        pages_with_text = 0
        # P1-4：按阅读顺序登记全文档表格（index 在汇合后统一补齐）
        all_tables: list[dict] = []
        # P1-5：页码行号区间 [[start_line, end_line, page], ...]，chunk 按段落起始行查表
        page_spans: list[list[int]] = []
        for page in pdf:
            page_start = global_line
            page_text = _extract_page_with_tables(page, all_tables)
            page_lines = page_text.split("\n")
            for heading in _pdf_page_headings(page, body_size):
                try:
                    line_in_page = page_lines.index(heading)
                except ValueError:
                    continue
                sections.append((heading, global_line + line_in_page))
            global_line += len(page_lines)
            parts.append(page_text)
            page_spans.append([page_start, global_line - 1, page.number + 1])
            total_chars += len(page_text)
            if page_text.strip():
                pages_with_text += 1
        metadata = {
            # 单页文档给出页码；多页文档页码按段归属留待 W2 细化（chunk 级 page 标注）
            "page": 1 if pdf.page_count == 1 else None,
            "author": (pdf.metadata or {}).get("author") or None,
            # P0-1：解析探针（garbled_ratio 由 _build_document 统一补齐）
            "text_stats": {
                "total_chars": total_chars,
                "page_count": pdf.page_count,
                "pages_with_text": pages_with_text,
            },
        }
        if sections:
            metadata["sections"] = sections   # T5：字号启发式章节标题
        if pdf.page_count > 1 and page_spans:
            metadata["page_spans"] = page_spans   # P1-5：多页 PDF 页码行号区间（单页已由 page=1 覆盖）
        if all_tables:
            for i, t in enumerate(all_tables):
                t["index"] = i
            metadata["tables"] = all_tables   # P1-4：顺序对应正文第 N 个 [TABLE] 块
        date = _normalize_pdf_date((pdf.metadata or {}).get("creationDate"))
        if date:
            metadata["date"] = date
        return "\n".join(parts), metadata
    finally:
        pdf.close()


def _extract_page_with_tables(page, tables_out: list | None = None) -> str:
    """按 y 坐标把页内文本块与表格（Markdown 化）合并为阅读顺序的文本。

    表格块标记为 [TABLE]...[/TABLE]，与正文段落用空行隔开，供后续分块/清洗识别；
    P1-4：tables_out 非空时按阅读顺序登记 {header, rows, page}（index 由调用方统一补）。
    """
    blocks: list[tuple[float, str]] = []
    raw_tables: list[tuple[float, list[list[str]]]] = []
    # 文本块（dict 模式带坐标）
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        text = " ".join(
            span["text"]
            for line in block.get("lines", [])
            for span in line.get("spans", [])
        ).strip()
        if text:
            blocks.append((block["bbox"][1], text))  # y0
    # 表格块
    try:
        tables = page.find_tables()
        for table in tables:
            rows = _table_rows(table)
            md = _rows_to_markdown(rows)
            if md:
                blocks.append((table.bbox[1], f"[TABLE]\n{md}\n[/TABLE]"))
                raw_tables.append((table.bbox[1], rows))
    except Exception:
        pass  # 表格抽取失败不阻断正文
    blocks.sort(key=lambda b: b[0])
    if tables_out is not None:
        for _y, rows in sorted(raw_tables, key=lambda t: t[0]):
            tables_out.append({"header": rows[0], "rows": rows[1:],
                               "page": getattr(page, "number", None)})
    return "\n".join(text for _, text in blocks)


def _table_rows(table) -> list[list[str]]:
    """表格 → 清洗后的二维行（P1-4：markdown 与 tables JSON 共用同一份行数据）。"""
    rows = table.extract() or []
    rows = [[(c or "").replace("\n", " ").strip() for c in row] for row in rows]
    return [row for row in rows if any(row)]


def _rows_to_markdown(rows: list[list[str]]) -> str:
    """行数据 → Markdown 管道表格（首行表头 + 分隔行）。"""
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |",
             "|" + "|".join("---" for _ in rows[0]) + "|"]
    for row in rows[1:]:
        while len(row) < len(rows[0]):
            row.append("")
        lines.append("| " + " | ".join(row[:len(rows[0])]) + " |")
    return "\n".join(lines)


def _table_to_markdown(table) -> str:
    """表格 → Markdown 管道表格（兼容旧调用；P1-4 起内部走 _table_rows/_rows_to_markdown）。"""
    return _rows_to_markdown(_table_rows(table))


def _parse_docx(raw: bytes) -> tuple[str, dict]:
    """Word：python-docx 按 body 文档流保序解析（段落 + [TABLE] 表格块，P1-4）。

    修复旧实现的两处缺陷：表格无标记且固定追加在段落末尾——现按 body 顺序
    遍历 w:p/w:tbl，表格在原位输出 [TABLE] 块并登记 metadata["tables"]。
    """
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    try:
        doc = DocxDocument(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"Word 解析失败（文件损坏或非 docx）: {exc}") from exc

    lines: list[str] = []
    sections: list[tuple[str, int]] = []
    tables_meta: list[dict] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = DocxParagraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            if style_name.startswith("Heading") or style_name.startswith("标题"):
                sections.append((text, len(lines)))   # T5：Heading 样式章节标题（行号=当前行）
            lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = DocxTable(child, doc)
            rows = [[c.text.strip() for c in row.cells] for row in table.rows]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            md = _rows_to_markdown(rows)
            lines.append("[TABLE]")
            lines.extend(md.split("\n"))
            lines.append("[/TABLE]")
            tables_meta.append({"header": rows[0], "rows": rows[1:], "page": None})

    metadata: dict = {"page": None, "author": None}
    if sections:
        metadata["sections"] = sections       # T5：Heading 样式章节标题
    if tables_meta:
        for i, t in enumerate(tables_meta):
            t["index"] = i
        metadata["tables"] = tables_meta      # P1-4：顺序对应正文第 N 个 [TABLE] 块
    try:
        created = doc.core_properties.created
        if created:
            metadata["date"] = str(created.date())
    except Exception:
        pass
    try:
        metadata["author"] = doc.core_properties.author or None
    except Exception:
        metadata["author"] = None
    return "\n".join(lines), metadata


def _parse_xlsx(raw: bytes) -> tuple[str, dict]:
    """Excel：每 sheet 每行判为一段文本；metadata 记录 sheet 名。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Excel 解析失败（文件损坏或非 xlsx）: {exc}") from exc
    try:
        lines: list[str] = []
        sheets: list[str] = []
        for ws in wb.worksheets:
            sheets.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append("\t".join(vals))
        return "\n".join(lines), {"page": None, "author": None, "sheets": sheets}
    finally:
        wb.close()


def _parse_pptx(raw: bytes) -> tuple[str, dict]:
    """PPT：python-pptx 按幻灯片/文本框抽取。"""
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"PPT 解析失败（文件损坏或非 pptx）: {exc}") from exc
    lines: list[str] = []
    page_spans: list[list[int]] = []
    for slide_no, slide in enumerate(prs.slides, 1):
        start = len(lines)
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
        if len(lines) > start:
            page_spans.append([start, len(lines) - 1, slide_no])
    return "\n".join(lines), {"page": None, "author": None, "page_spans": page_spans}

def _parse_html(raw: bytes) -> tuple[str, dict]:
    """HTML：标准库 html.parser 提取正文文本（去 script/style/注释），零新增依赖。"""
    from html.parser import HTMLParser

    class _TextExtractor(HTMLParser):
        SKIP_TAGS = {"script", "style", "noscript", "template"}

        def __init__(self):
            super().__init__()
            self.out: list[str] = []
            self._skip_depth = 0
            self._block = False

        def handle_starttag(self, tag, attrs):
            if tag in self.SKIP_TAGS:
                self._skip_depth += 1
            if self._skip_depth == 0 and tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table"):
                self._block = True

        def handle_endtag(self, tag):
            if tag in self.SKIP_TAGS:
                self._skip_depth = max(0, self._skip_depth - 1)
            if self._skip_depth == 0 and tag in ("p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table"):
                self._block = True

        def handle_data(self, data):
            if self._skip_depth == 0 and data.strip():
                self.out.append(data.strip())

        def close(self):
            super().close()
            return "\n".join(self.out)

    try:
        text = raw.decode("utf-8", errors="replace")
        parser = _TextExtractor()
        parser.feed(text)
        body = parser.close()
    except Exception as exc:
        raise ParseError(f"HTML 解析失败: {exc}") from exc
    return body, {"page": None, "author": None}


def _parse_eml(raw: bytes) -> tuple[str, dict]:
    """邮件 .eml：标准库 email 解析，正文取 text/plain（无则 text/html 转文本）。"""
    import email
    from email import policy

    try:
        msg = email.message_from_bytes(raw, policy=policy.default)
    except Exception as exc:
        raise ParseError(f"邮件解析失败（非 .eml）: {exc}") from exc
    subject = str(msg.get("subject") or "").strip()
    sender = str(msg.get("from") or "").strip()
    parts: list[str] = []
    for part in msg.walk():
        if part.get_content_maintype() == "multipart":
            continue
        ctype = part.get_content_type()
        if ctype == "text/plain":
            payload = part.get_payload(decode=True)
            if payload:
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace").strip())
        elif ctype == "text/html" and not parts:
            payload = part.get_payload(decode=True)
            if payload:
                html_text, _ = _parse_html(payload)
                if html_text:
                    parts.append(html_text)
    body = "\n\n".join(p for p in parts if p)
    if subject:
        body = f"主题：{subject}\n\n{body}".strip()
    return body, {"page": None, "author": sender or None}


def _parse_ocr(raw: bytes, ext: str) -> tuple[str, dict]:
    """图片 OCR 预留接口（DEVELOPMENT.md 13.10 P3：parser.ocr.enabled=false 默认关）。

    未接入引擎时抛 ParseError（明确不可用，不静默返回空文档）；
    接入引擎（如 paddleocr/pytesseract）只需替换本函数内部实现。
    """
    raise ParseError(f"OCR 未启用（图片 {ext} 暂不支持，需开启 parser.ocr.enabled 并接入引擎）")


PARSERS: dict[str, callable] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".pptx": _parse_pptx,
    ".md": _parse_text,
    ".txt": _parse_text,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".eml": _parse_eml,
    ".msg": _parse_eml,
}
