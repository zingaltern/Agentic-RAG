"""Office 解析（DEVELOPMENT.md 5.1 / 16.2）：docx/xlsx/pptx + 旧版 xls/doc/ppt。

- 现代格式：python-docx / openpyxl / python-pptx；
- 旧版 .xls：xlrd（纯 Python）；
- 旧版 .doc/.ppt：经 LibreOffice 转换后复用 docx/pptx 解析（见 office_convert）。
"""
from __future__ import annotations

import io

from app.adapters.parser.office_convert import _convert_legacy_office
from app.adapters.parser.quality import _rows_to_markdown
from app.core.errors import ParseError


def _parse_docx(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """Word：python-docx 按 body 文档流保序解析（段落 + [TABLE] 表格块，P1-4）。

    修复旧实现的两处缺陷：表格无标记且固定追加在段落末尾——现按 body 顺序
    遍历 w:p/w:tbl，表格在原位输出 [TABLE] 块并登记 metadata["tables"]。
    """
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    try:
        doc = DocxDocument(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"Word 解析失败（文件损坏或非 docx）: {exc}") from exc

    lines: list[str] = []
    sections: list[tuple[str, int]] = []
    tables_meta: list[dict] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = DocxParagraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            if style_name.startswith("Heading") or style_name.startswith("标题"):
                sections.append((text, len(lines)))   # T5：Heading 样式章节标题（行号=当前行）
            lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = DocxTable(child, doc)
            rows = [[c.text.strip() for c in row.cells] for row in table.rows]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            md = _rows_to_markdown(rows)
            lines.append("[TABLE]")
            lines.extend(md.split("\n"))
            lines.append("[/TABLE]")
            tables_meta.append({"header": rows[0], "rows": rows[1:], "page": None})

    metadata: dict = {"page": None, "author": None}
    if sections:
        metadata["sections"] = sections       # T5：Heading 样式章节标题
    if tables_meta:
        for i, t in enumerate(tables_meta):
            t["index"] = i
        metadata["tables"] = tables_meta      # P1-4：顺序对应正文第 N 个 [TABLE] 块
    try:
        created = doc.core_properties.created
        if created:
            metadata["date"] = str(created.date())
    except Exception:
        pass
    try:
        metadata["author"] = doc.core_properties.author or None
    except Exception:
        metadata["author"] = None
    return "\n".join(lines), metadata

def _parse_xlsx(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """Excel：每 sheet 每行判为一段文本；metadata 记录 sheet 名。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Excel 解析失败（文件损坏或非 xlsx）: {exc}") from exc
    try:
        lines: list[str] = []
        sheets: list[str] = []
        for ws in wb.worksheets:
            sheets.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append("\t".join(vals))
        return "\n".join(lines), {"page": None, "author": None, "sheets": sheets}
    finally:
        wb.close()

def _parse_pptx(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """PPT：python-pptx 按幻灯片/文本框抽取。"""
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as exc:
        raise ParseError(f"PPT 解析失败（文件损坏或非 pptx）: {exc}") from exc
    lines: list[str] = []
    page_spans: list[list[int]] = []
    for slide_no, slide in enumerate(prs.slides, 1):
        start = len(lines)
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
        if len(lines) > start:
            page_spans.append([start, len(lines) - 1, slide_no])
    return "\n".join(lines), {"page": None, "author": None, "page_spans": page_spans}

def _cell_text(value: object) -> str:
    """xls 单元格 → 文本（xlrd 数字统一为 float，整数去掉 `.0`）。"""
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def _parse_xls(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """旧版 Excel .xls（16.2，xlrd）：逐 sheet 逐行取值，与 xlsx 同款「制表符连接」文本。"""
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:
        raise ParseError(f"Excel(.xls) 解析失败（文件损坏或非 xls）: {exc}") from exc
    lines: list[str] = []
    sheets: list[str] = []
    for sheet in book.sheets():
        sheets.append(sheet.name)
        for row_index in range(sheet.nrows):
            vals = [_cell_text(sheet.cell_value(row_index, col)) for col in range(sheet.ncols)]
            vals = [v for v in vals if v]
            if vals:
                lines.append("\t".join(vals))
    return "\n".join(lines), {"page": None, "author": None, "sheets": sheets}

def _parse_legacy_office(raw: bytes, source_ext: str, target_ext: str,
                         config: object | None) -> tuple[str, dict]:
    """转换 → 复用既有 docx/pptx 解析 → 补溯源 metadata。"""
    converted = _convert_legacy_office(raw, source_ext, target_ext, config)
    if target_ext == ".docx":
        text, meta = _parse_docx(converted, config)
    else:
        text, meta = _parse_pptx(converted, config)
    meta = {**(meta or {}), "parser": "soffice_convert",
            "converted_from": source_ext.lstrip(".")}
    return text, meta

def _parse_doc(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """旧版 Word .doc（16.2）：LibreOffice 转 .docx 后复用 python-docx 解析。"""
    return _parse_legacy_office(raw, ".doc", ".docx", config)

def _parse_ppt(raw: bytes, config: object | None = None) -> tuple[str, dict]:
    """旧版 PPT .ppt（16.2）：LibreOffice 转 .pptx 后复用 python-pptx 解析。"""
    return _parse_legacy_office(raw, ".ppt", ".pptx", config)
